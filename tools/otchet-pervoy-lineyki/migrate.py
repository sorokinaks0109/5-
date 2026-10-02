"""Перенос в новую схему (HTML + JSON) из старого общего отчёта и свежей карты УПЦ.

Запуск: python3 migrate.py старый_отчёт.xlsx карта_УПЦ.xlsx "папка/Отчёт первой линейки"
Создаёт папку «Данные»: реестр.json, руководители/<Фамилия>.json, вложения/<Фамилия>/Приложения.xlsx.
"""
import copy
import datetime as dt
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

SRC, CARD, OUT = sys.argv[1:4]
DATA = os.path.join(OUT, 'Данные')
MAX_COLS = 12

# Кто отчитывается за показатель, если лидер не из первой линейки (решение заказчика).
REPORTER_BY_LEADER = {'Топинская': 'Арикулова', 'Усманов': 'Дрыков', 'Бруховский': None, 'Губаревич': None}
EXTRA_PEOPLE = [('Лазар А.А.', False)]  # заполняет отчёт, на совещании не докладывает

DONE_RE = re.compile(r'(^|[^а-яё])(выполнен|проведен|исполнен|завершен|снят|закрыт)(о|а|ы)?(?![а-яё])', re.I)
NOT_DONE_RE = re.compile(r'(^|[^а-яё])не\s+(выполнен|проведен|исполнен|завершен|снят|закрыт)', re.I)


def is_done(t):
    t = str(t or '')
    return bool(DONE_RE.search(t)) and not NOT_DONE_RE.search(t)


def clean(v):
    if isinstance(v, str):
        v = v.replace('\xa0', ' ').strip()
        return v or None
    return v


def iso(v):
    if isinstance(v, dt.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, str):
        m = re.fullmatch(r'(\d{1,2})\.(\d{1,2})\.(\d{4})', v.strip())
        if m:
            return f'{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}'
        return v.strip()
    if v is None:
        return ''
    return str(v)


def text(v):
    v = clean(v)
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime('%d.%m.%Y')
    return str(v)


def surname(s):
    return re.split(r'[\s.]+', str(s or '').strip())[0].lower().replace('ё', 'е')


wb = openpyxl.load_workbook(SRC, data_only=True)
wb_st = openpyxl.load_workbook(SRC)
card = openpyxl.load_workbook(CARD, data_only=True).active

FIRST = [n for n in wb.sheetnames if n not in ('Как заполнять', 'Мемо', 'План мероприятий для УПЦ', 'УПЦ справочно')]

# ---------- карта УПЦ ----------
head_row, cols, card_date = None, {}, ''
for r in range(1, 30):
    vals = [text(card.cell(r, c).value) for c in range(1, 20)]
    for i, v in enumerate(vals):
        if re.search(r'дата выгрузки', v, re.I):
            for cand in [vals[i + 1] if i + 1 < len(vals) else '', text(card.cell(r + 1, i + 1).value), text(card.cell(r + 1, i + 2).value)]:
                if re.search(r'\d{2}\.\d{2}\.\d{4}', cand):
                    card_date = iso(re.search(r'\d{2}\.\d{2}\.\d{4}', cand)[0])
    if 'Идентификатор' in vals and 'Показатель' in vals:
        head_row = r
        cols = {v: i + 1 for i, v in enumerate(vals) if v}
        break
kpis = []
for r in range(head_row + 1, card.max_row + 1):
    g = lambda name: text(card.cell(r, cols[name]).value) if name in cols else ''
    if not g('Идентификатор') or not g('Показатель'):
        continue
    kpis.append({
        'id': g('Идентификатор'), 'top3': g('ТОП 3').lower() == 'да', 'name': g('Показатель'), 'goal': g('Цель'),
        'stretch': g('Напряженная цель'), 'fact': g('Факт (на дату мониторинга)'), 'forecast': g('Прогноз (на 31.12)'),
        'grade': g('Оценка прогноза'), 'team': g('Команда ответственных'), 'leader': g('Лидер команды'),
        'unit': g('Ед. Изм'), 'type': g('Тип'),
    })

# ---------- люди ----------
people = []
for name, on in [(n, True) for n in FIRST] + EXTRA_PEOPLE:
    sn = surname(name)
    full = ''
    for k in kpis:
        for part in re.split(r'[,()]', k['team'] + ',' + k['leader']):
            part = part.strip()
            if surname(part) == sn and len(part.split()) == 3:
                full = part
    slug = name.split()[0]
    people.append({'id': slug, 'fio': name.strip(), 'full': full, 'slug': slug, 'email': '', 'onMeeting': on})
by_surname = {surname(p['fio']): p for p in people}


for k in kpis:
    ls = surname(k['leader'])
    key = next((x for x in REPORTER_BY_LEADER if surname(x) == ls), None)
    if key is not None:
        target = REPORTER_BY_LEADER[key]
        k['reporter'] = by_surname[surname(target)]['id'] if target else None
    else:
        k['reporter'] = by_surname[ls]['id'] if ls in by_surname else None

# ---------- поручения ----------
ws = wb['Мемо']
memo = []
for r in range(5, ws.max_row + 1):
    mid = text(ws.cell(r, 1).value)
    if not mid:
        continue
    resp_name = text(ws.cell(r, 4).value)
    rp = by_surname.get(surname(resp_name))
    co_names = [x.strip() for x in re.split(r'[,;]', text(ws.cell(r, 5).value)) if x.strip()]
    co_ids = [by_surname[surname(x)]['id'] for x in co_names if surname(x) in by_surname]
    co_other = [x for x in co_names if surname(x) not in by_surname]
    num = int(re.search(r'\d+', mid)[0])
    item = {
        'id': mid, 'num': num, 'date': iso(ws.cell(r, 2).value) if text(ws.cell(r, 2).value) != 'ранее' else '',
        'text': text(ws.cell(r, 3).value) + (('\nСоисп.: ' + ', '.join(co_other)) if co_other else ''),
        'resp': rp['id'] if rp else '', 'respText': '' if rp else resp_name, 'co': co_ids, 'due': iso(ws.cell(r, 6).value),
    }
    mark = text(ws.cell(r, 8).value)
    if is_done(mark):
        item['closed'] = {'date': '', 'flag': 'done', 'text': mark}
    memo.append(item)

# ---------- мероприятия УПЦ ----------
def core(s):
    s = re.sub(r'(ГПН-ГПН_Снаб-)+', '', str(s or ''))
    return re.sub(r'\s+', ' ', s).strip().lower()


ws = wb['План мероприятий для УПЦ']
events, cur_kpi, dropped, manual = [], None, [], {}
for r in range(7, ws.max_row + 1):
    a, b, c, d, e, f, g, h = [clean(ws.cell(r, k).value) for k in range(1, 9)]
    if not h and a and not b and isinstance(a, str) and not re.fullmatch(r'\d+\.?', a.strip()):
        if a.lower().startswith('блок'):
            continue
        heading = core(a)
        # «Качество планирования_Сибирь» ↔ «Качество планирования_Восточная сибирь»: сравниваем часть до «_»
        found = [k for k in kpis if core(k['name']).split('_')[0][:60] in heading]
        cur_kpi = max(found, key=lambda k: len(core(k['name']).split('_')[0][:60])) if found else None
        if not cur_kpi:
            dropped.append(a.strip())
        continue
    if not h or not cur_kpi:
        if h:
            dropped.append(f'  {h} {text(b)[:50]}')
        continue
    speaker = by_surname.get(surname(g))['id'] if g and surname(g) in by_surname else None
    if not speaker:
        speaker = cur_kpi['reporter']
    if not speaker:
        ex = [x for x in re.split(r'[,;]', text(c)) if x.strip()]
        if len(ex) == 1 and surname(ex[0]) in by_surname:
            speaker = by_surname[surname(ex[0])]['id']
    ev = {'id': h, 'kpi': cur_kpi['id'], 'text': text(b), 'executors': text(c), 'due': iso(d), 'speaker': speaker}
    if f and not g:
        manual[h] = text(f)
    events.append(ev)

# ---------- личные отчёты ----------
def flag_of(t):
    return 'done' if is_done(t) else ('work' if t else '')


def free_row(vals):
    a, b, c, dd, e, f = (vals + [None] * 6)[:6]
    extra = [text(x) for x in vals[6:] if clean(x) is not None]
    due = ''
    if isinstance(e, dt.datetime) or (isinstance(e, str) and re.fullmatch(r'\d{1,2}\.\d{1,2}\.\d{4}', e.strip())):
        due = iso(e)
        pct = f if isinstance(f, (int, float)) and 0 <= f <= 100 else ''
        if pct == '' and clean(f) is not None:
            extra.insert(0, text(f))
    else:
        pct = ''
        extra = [text(x) for x in (e, f) if clean(x) is not None] + extra
    body = text(c) + ((' | ' + ' | '.join(extra)) if extra else '')
    return {'kind': 'row', 'task': text(b), 'text': body, 'resp': text(dd), 'due': due, 'flag': flag_of(text(c)) if text(c) else '', 'pct': pct}


os.makedirs(os.path.join(DATA, 'руководители'), exist_ok=True)
own = {p['id']: {'personId': p['id'], 'fio': p['fio'], 'reportDate': '', 'memo': {}, 'events': {}, 'kpi': {}, 'ros': [], 'proj': [], 'attachments': []} for p in people}
memo_resp = {m['id']: m['resp'] for m in memo}
ev_speaker = {e['id']: e['speaker'] for e in events}

for name in FIRST:
    ws = wb[name]
    ws_st = wb_st[name]
    pid = by_surname[surname(name)]['id']
    rows = {r: [clean(ws.cell(r, c).value) for c in range(1, MAX_COLS + 1)] for r in range(1, ws.max_row + 1)}
    title = lambda v, w: isinstance(v[0], str) and v[0].strip().upper().startswith(w)
    memo_start = next(r for r, v in rows.items() if title(v, 'ПОРУЧЕНИЯ ПО МЕМО'))
    upc_start = next(r for r, v in rows.items() if title(v, 'МЕРОПРИЯТИЯ УПЦ'))
    cur = None
    for r in range(2, memo_start):
        v = [None if x == '•' else x for x in rows[r]]
        while v and v[-1] is None:
            v = v[:-1]
        if not v or v[0] in ('№ ', '№') or (len([x for x in v if x is not None]) == 1 and str(v[0]).strip().upper() == name.upper()):
            continue
        is_head = isinstance(v[0], str) and not re.fullmatch(r'\s*\d+\.?\s*', v[0]) and all(x is None for x in (v + [None] * 5)[1:5])
        if is_head:
            hd = v[0].strip()
            if hd.upper() == 'РОС':
                cur = 'ros'
            elif hd.lower().startswith('текущие проекты'):
                cur = 'proj'
            else:
                own[pid][cur or 'proj'].append({'kind': 'head', 'task': hd})
            continue
        if all(x is None for x in v[1:]):
            continue
        own[pid][cur or 'proj'].append(free_row(v))
    for start, kind in ((memo_start, 'memo'), (upc_start, 'events')):
        for r in range(start + 2, ws.max_row + 1):
            v = rows[r]
            key = v[0].strip() if isinstance(v[0], str) else ''
            if not re.fullmatch(r'[МУ]-\d+', key):
                if key and r > start + 2:
                    break
                continue
            st = text(v[2])
            if not st:
                continue
            owner = (memo_resp if kind == 'memo' else ev_speaker).get(key)
            if owner == pid or (kind == 'events' and owner and owner not in own):
                own[pid][kind][key] = {'flag': flag_of(st), 'text': st}
            elif kind == 'events' and owner:
                own[owner]['events'].setdefault(key, {'flag': flag_of(st), 'text': st})
    # приложения: всё ниже блока УПЦ, что не строка УПЦ
    app_start = next((r for r in range(upc_start + 2, ws.max_row + 1)
                      if any(x is not None for x in rows[r]) and not (isinstance(rows[r][0], str) and re.fullmatch(r'У-\d+', rows[r][0].strip()))), None)
    if app_start:
        out = openpyxl.Workbook()
        app = out.active
        app.title = 'Приложения'
        last = max(r for r in range(app_start, ws.max_row + 1) if any(x is not None for x in rows[r]))
        for r in range(app_start, last + 1):
            for c in range(1, MAX_COLS + 1):
                src = ws_st.cell(r, c)
                dst = app.cell(r - app_start + 1, c, clean(ws.cell(r, c).value))
                if src.has_style:
                    dst.font, dst.fill, dst.border, dst.alignment = copy.copy(src.font), copy.copy(src.fill), copy.copy(src.border), copy.copy(src.alignment)
                    dst.number_format = src.number_format
        for mr in ws_st.merged_cells.ranges:
            if mr.min_row >= app_start and mr.max_row <= last and mr.max_col <= MAX_COLS:
                app.merge_cells(start_row=mr.min_row - app_start + 1, start_column=mr.min_col, end_row=mr.max_row - app_start + 1, end_column=mr.max_col)
        for c in range(1, MAX_COLS + 1):
            w = ws_st.column_dimensions[get_column_letter(c)].width
            if w:
                app.column_dimensions[get_column_letter(c)].width = w
        slug = by_surname[surname(name)]['slug']
        os.makedirs(os.path.join(DATA, 'вложения', slug), exist_ok=True)
        out.save(os.path.join(DATA, 'вложения', slug, 'Приложения.xlsx'))
        own[pid]['attachments'].append({'name': 'Приложения (перенесено из старого отчёта)', 'file': 'Приложения.xlsx', 'added': dt.date.today().isoformat()})

# статусы мероприятий без докладчика (ручные) и тех, у кого теперь появился докладчик
for e in events:
    t = manual.get(e['id'])
    if not t:
        continue
    if e['speaker']:
        own[e['speaker']]['events'].setdefault(e['id'], {'flag': flag_of(t), 'text': t})
    else:
        e['manual'] = {'flag': flag_of(t), 'text': t}

last_meeting = max((m['date'] for m in memo if re.fullmatch(r'\d{4}-\d{2}-\d{2}', m['date'] or '')), default='')
reg = {
    'version': 2,
    'settings': {'lastMeeting': last_meeting, 'cardDate': card_date, 'directorEmail': ''},
    'people': people, 'memo': memo, 'kpis': kpis, 'events': events,
}
with open(os.path.join(DATA, 'реестр.json'), 'w', encoding='utf-8') as f:
    json.dump(reg, f, ensure_ascii=False, indent=1)
for p in people:
    with open(os.path.join(DATA, 'руководители', p['slug'] + '.json'), 'w', encoding='utf-8') as f:
        json.dump(own[p['id']], f, ensure_ascii=False, indent=1)

print(f'Руководителей: {len(people)}, поручений: {len(memo)}, показателей: {len(kpis)}, мероприятий: {len(events)}, дата карты: {card_date}')
for p in people:
    o = own[p['id']]
    rep = [k['name'].replace('ГПН-ГПН_Снаб-', '')[:40] for k in kpis if k['reporter'] == p['id']]
    evs = sum(1 for e in events if e['speaker'] == p['id'])
    print(f"  {p['fio']}: РОС {len(o['ros'])}, проекты {len(o['proj'])}, статусов поручений {len(o['memo'])}, мероприятий {evs} (статусов {len(o['events'])}), отчитывается: {rep}")
print('Без отчитывающегося:', [k['name'].replace('ГПН-ГПН_Снаб-', '')[:50] for k in kpis if not k['reporter']])
print('Мероприятий без докладчика:', sum(1 for e in events if not e['speaker']))
print('Не перенесено (показателя нет в новой карте):')
for d in dropped:
    print('  ', d[:110])
