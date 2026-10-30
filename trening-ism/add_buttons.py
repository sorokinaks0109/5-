# -*- coding: utf-8 -*-
"""Добавляет в презентацию кнопки, которые открывают интерактивные упражнения.

Запуск: python3 add_buttons.py исходная.pptx Тренинг_ИСМ.pptx
"""
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

from build import EXERCISES

# упражнение 06 относится к заданию на слайде 31, а не к реестру на 32
SLIDES = {ex["slide"]: ex["file"] for ex in EXERCISES}

src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
W, H = prs.slide_width, prs.slide_height
bw, bh = Emu(2300000), Emu(560000)

for num, file in SLIDES.items():
    slide = prs.slides[num - 1]
    btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, W - bw - Emu(250000), H - bh - Emu(200000), bw, bh)
    btn.name = "Кнопка интерактив"
    btn.fill.solid()
    btn.fill.fore_color.rgb = RGBColor(0xE6, 0x59, 0x07)
    btn.line.fill.background()
    tf = btn.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "▶  ИНТЕРАКТИВ"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    btn.click_action.hyperlink.address = "uprazhneniya/" + file

prs.save(dst)
print(f"Кнопки добавлены на слайды: {sorted(SLIDES)}")
